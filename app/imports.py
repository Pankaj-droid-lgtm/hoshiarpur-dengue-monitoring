from datetime import datetime
from hashlib import sha256
from io import BytesIO

from flask import Blueprint, flash, redirect, render_template, request, url_for
from flask_login import login_required
from openpyxl import load_workbook

from .extensions import db
from .models import HistoricalBlockReport, HistoricalDengueCase, HistoricalFieldResponse, HistoricalStaffingAllocation, SourceDocument
from .security import validate_csrf
from .services.audit import log_change
from .services.permissions import require_management_access


imports_bp = Blueprint("imports", __name__, url_prefix="/imports")


@imports_bp.route("/historical", methods=["GET", "POST"])
@login_required
def historical_import():
    require_management_access()
    if request.method == "POST":
        validate_csrf()
        upload = request.files.get("source_file")
        if not upload or not upload.filename.lower().endswith(".xlsx"):
            flash("Select an Excel .xlsx source file.", "error")
            return redirect(url_for("imports.historical_import"))
        contents = upload.read()
        digest = sha256(contents).hexdigest()
        if SourceDocument.query.filter_by(sha256=digest).first():
            flash("This source document was already imported.", "error")
            return redirect(url_for("imports.historical_import"))
        try:
            workbook = load_workbook(BytesIO(contents), read_only=True, data_only=False)
        except Exception:
            flash("The uploaded file could not be read as an Excel workbook.", "error")
            return redirect(url_for("imports.historical_import"))
        source = SourceDocument(original_filename=upload.filename, sha256=digest, source_type=request.form.get("source_type", "other"))
        db.session.add(source)
        db.session.flush()
        imported = import_workbook(workbook, source)
        log_change("import_source", "source_document", source.id, after={"filename": source.original_filename, "rows": imported})
        db.session.commit()
        flash(f"Imported {imported} historical source rows. No operational records were created.", "success")
        return redirect(url_for("imports.historical_import"))
    return render_template("imports/historical.html", sources=SourceDocument.query.order_by(SourceDocument.imported_at.desc()).all())


def import_workbook(workbook, source: SourceDocument) -> int:
    filename = source.original_filename.casefold()
    if "dengue cases" in filename:
        return import_cases(workbook, source)
    if "breeding checkers" in filename:
        return import_staffing(workbook, source)
    return import_block_reporting(workbook, source)


def import_cases(workbook, source: SourceDocument) -> int:
    sheet = workbook["Sheet1"]
    count = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        if not row or not row[0]:
            continue
        values = list(row) + [None] * 10
        db.session.add(HistoricalDengueCase(source_document_id=source.id, source_row=row_number, source_serial=string(values[0]), patient_name=string(values[1]), age_years=integer(values[2]), sex_gender_raw=string(values[3]), contact_number_raw=string(values[4]), address_raw=string(values[5]), rural_urban_raw=string(values[6]), block_raw=string(values[7]), town_raw=string(values[8]), testing_date_raw=string(values[9])))
        count += 1
    return count


def import_staffing(workbook, source: SourceDocument) -> int:
    sheet = workbook["Sheet1"]
    count = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        if not row or str(row[0]).strip().casefold() == "total":
            continue
        values = list(row) + [None] * 6
        db.session.add(HistoricalStaffingAllocation(source_document_id=source.id, source_row=row_number, rural_area_raw=string(values[1]), urban_area_raw=string(values[2]), breeding_checker_count=integer(values[3]), mphw_male_count=integer(values[4]), team_count=integer(values[5])))
        count += 1
    return count


def import_block_reporting(workbook, source: SourceDocument) -> int:
    count = 0
    for sheet in workbook.worksheets:
        if sheet.title in {"Instructions", "Weekly Block Summary"}:
            continue
        for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
            if not row or not row[0] or str(row[0]).strip().casefold() == "district total":
                continue
            values = {f"column_{index + 1}": raw_value(value) for index, value in enumerate(row)}
            if sheet.title == "Field Visit & Case Response":
                db.session.add(HistoricalFieldResponse(source_document_id=source.id, source_sheet=sheet.title, source_row=row_number, block_raw=string(row[0]), locality_raw=string(row[1]), visit_date_raw=string(row[2]), values=values))
            else:
                db.session.add(HistoricalBlockReport(source_document_id=source.id, source_sheet=sheet.title, source_row=row_number, block_raw=string(row[0]), reporting_period=None, values=values))
            count += 1
    return count


def string(value):
    return str(value).strip() if value is not None and str(value).strip() else None


def integer(value):
    try:
        return int(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def raw_value(value):
    return value.isoformat() if isinstance(value, datetime) else value
